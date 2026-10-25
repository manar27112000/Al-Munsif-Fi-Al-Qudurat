# -*- coding: utf-8 -*-
"""اختبار مسار التصدير كامل بعميل تليجرام مزيّف (من غير إنترنت)."""
import asyncio, json, os, shutil, sys, tempfile, zipfile
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "app", "src", "main", "python"))
from datetime import datetime, timezone, timedelta
from telethon import types
import bridge

jpeg = None
with zipfile.ZipFile(os.path.join(HERE, "fixture.zip")) as z:
    for n in z.namelist():
        if n.lower().endswith((".jpg", ".jpeg")):
            jpeg = z.read(n); break
assert jpeg, "no jpeg in fixture"

base = datetime(2026, 10, 20, 10, 0, tzinfo=timezone.utc)

def photo_media():
    ph = types.Photo(id=1, access_hash=1, file_reference=b"x", date=base, sizes=[], dc_id=1)
    return types.MessageMediaPhoto(photo=ph)

def mk(i, minutes, text, user=5, media=None):
    return types.Message(id=i, peer_id=types.PeerChannel(1), date=base + timedelta(minutes=minutes),
                         message=text, from_id=types.PeerUser(user), media=media)

MSGS = [
    mk(1, -600, "قديمة جدًا"),
    mk(2, 0, "1) ما ناتج ٢+٢ ؟\n❇️ الرسالة موجهة من حساب : AL\n🆔 7990637609"),
    mk(3, 0.5, "", media=photo_media()),
    mk(4, 10, "2) سؤال تاني (@Giyuu111)"),
    mk(5, 20, "1) ما ناتج ٢+٢ ؟\n🆔 7990637609"),   # مكرر (نص + صورة)
    mk(7, 20.2, "", media=photo_media()),
    types.MessageService(id=6, peer_id=types.PeerChannel(1), date=base + timedelta(minutes=21), action=types.MessageActionPinMessage()),
]

class FakeClient:
    async def get_input_entity(self, cid): return cid
    async def iter_dialogs(self): 
        if False: yield
    def iter_messages(self, entity, offset_date=None):
        async def gen():
            for m in sorted(MSGS, key=lambda x: x.date, reverse=True):
                if offset_date and m.date >= offset_date: continue
                yield m
        return gen()
    async def download_media(self, m, file=None):
        open(file, "wb").write(jpeg); return file

bridge._client = FakeClient()
bridge._state["dir"] = tempfile.mkdtemp()
out = tempfile.mkdtemp()
from_ts = int((base - timedelta(hours=1)).timestamp())
to_ts = int((base + timedelta(hours=1)).timestamp())
r = json.loads(bridge.export(-100123, "اختبار", from_ts, to_ts, 60, out, "اختبار_٢٠٢٦"))
print(r)
assert r["ok"], r
from pptx import Presentation
p = Presentation(r["path"])
for i, s in enumerate(p.slides, 1):
    print(i, [(sh.shape_type, sh.text_frame.text if sh.has_text_frame else "") for sh in s.shapes][:2])
assert r["slides"] == 2 and r["removed"] == 1, r
txt = " ".join(sh.text_frame.text for s in p.slides for sh in s.shapes if sh.has_text_frame)
assert "7990637609" not in txt and "Giyuu111" not in txt and "الرسالة موجهة" not in txt
assert not os.path.exists(os.path.join(out, [d for d in os.listdir(out) if d.startswith("work_")][0])) if any(d.startswith("work_") for d in os.listdir(out)) else True
print("OK")
